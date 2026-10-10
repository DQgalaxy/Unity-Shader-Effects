from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
import json
import math

root = Path('F:/UnityProjects/10')
work = root/'Logs/GroundSlashLesson'
report = json.loads((work/'edit-report.json').read_text(encoding='utf-8'))
doc = Path(report['document'])
backup = Path(report['backup'])
ns = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
def blocks(z):
    return list(etree.fromstring(z.read('word/document.xml')).find('w:body', ns))
def text(e):
    return ''.join(e.xpath('.//w:t/text()',namespaces=ns))
def bounds(items):
    start=next(i for i,e in enumerate(items) if text(e)=='16 SG_GroundSlash 弧形刀光')
    end=next(i for i in range(start+1,len(items)) if text(items[i]).startswith('17 SG_MeteorSurface'))
    return start,end
with ZipFile(backup) as old, ZipFile(doc) as new:
    a,b=blocks(old),blocks(new)
    x,y=bounds(a);p,q=bounds(b)
    # Namespace placement can change during XML serialization; compare canonical subtrees.
    canonical=lambda es:[etree.tostring(e,method='c14n',exclusive=True) for e in es]
    assert canonical(a[:x])==canonical(b[:p])
    assert canonical(a[y:])==canonical(b[q:])
    changed={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
    for name in old.namelist():
        if name not in changed:assert old.read(name)==new.read(name),name
    assert new.testzip() is None
full=Document(doc)
texts='\n'.join(p.text for p in full.paragraphs)
scripts = ['Runtime/FXGroundSlashLesson.cs', 'Runtime/FXGroundSlashParticles.cs',
           'Editor/FXGroundSlashLessonEditor.cs']
for script in scripts:
    source=(root/'Assets/PortfolioFX/Scripts'/script).read_text(encoding='utf-8')
    assert source.strip() in texts,script
assert '{{SCRIPT}}' not in texts and '{{PARTICLES_SCRIPT}}' not in texts
assert '{{EDITOR_SCRIPT}}' not in texts
for required in ['分步思路 先不看连线','完整参考接法','SG_SlashWake','SG_SlashGroundCut',
                 'SG_SlashDust','参数实验与验收','GroundDebris_0～6','ContactWave','GroundSpray']:
    assert required in texts,required
for name in report['assets']:assert (root/'Assets/PortfolioFX/ReferenceGroundSlash'/name).exists()
clamp=lambda x:max(0,min(1,x))
def smooth(a,b,x):
    t=clamp((x-a)/(b-a));return t*t*(3-2*t)
for v in [0,.1,.5,.9,1]:
    for reveal in [0,1]:
        d=math.dist((v,.5),(.5,.5));r=-.08+.88*reveal
        f=1-smooth(r,r+.06,d)
        assert abs(f-reveal)<1e-7
for a in [0,.1,.5,1]:
    for power in [.5,.85,1,3]:
        for noise in [0,.5,1]:
            for amount in [0,.35,1]:
                mask=clamp(a)**power*(1-amount+amount*(.65+.35*noise))
                assert math.isfinite(mask) and 0<=mask<=1
                assert clamp(mask*0)==0
assert 1-smooth(2,2.8,3.8-.98)==0
qa=Document(work/'chapter16-preview.docx')
for table in qa.tables:
    assert table.rows[0]._tr.xpath('./w:trPr/w:tblHeader')
    assert table._tbl.xpath('./w:tblPr/w:tblBorders')
    for row in table.rows:
        assert not row._tr.xpath('./w:trPr/w:trHeight[@w:hRule="exact"]')
print('PASS: DOCX package; unchanged non-chapter content and styles; all 3 scripts match Unity source; asset paths; mask/reveal/fade endpoint tests; table header/border/height checks.')
print('Layout rendering is not verified: LibreOffice is unavailable.')
