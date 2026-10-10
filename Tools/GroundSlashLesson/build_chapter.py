from pathlib import Path
from copy import deepcopy
from datetime import datetime
from zipfile import ZipFile, ZIP_DEFLATED
from hashlib import sha256
import json
import re
import shutil
import uuid

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from lxml import etree

ROOT = Path('F:/UnityProjects/10')
WORK = ROOT/'Logs/GroundSlashLesson'
DOC = ROOT/'工程10_ShaderGraph特效原理与实操教程.docx'
WORK.mkdir(parents=True, exist_ok=True)
BACKUP = ROOT/'RepairBackups'/('GroundSlashChapter16_'+datetime.now().strftime('%Y%m%d_%H%M%S'))
BACKUP.mkdir(parents=True)
shutil.copy2(DOC, BACKUP/DOC.name)

# Prepare copies only; never change the source packages or active shaders.
kit = ROOT/'Assets/PortfolioFX/ReferenceGroundSlash'
kit.mkdir(exist_ok=True)
kenney = Path('F:/UnityProjects/OthersDemo/kenney-particle-pack-master/addons/kenney_particle_pack')
magic = ROOT/'Assets/UnityTechnologies/ParticlePack/EffectExamples/Magic Effects/Textures'
sources = {
    'T_SlashBlade.png': kenney/'slash_03.png',
    'T_SlashCrescent.png': kenney/'slash_02.png',
    'T_SlashFlash.png': kenney/'star_01.png',
    'T_SlashTrail.png': ROOT/'Assets/Trail/Textures/Trail_02.png',
    'T_SlashCrack.tif': magic/'EarthCrackEmissive.tif',
    'T_SlashDust.png': magic/'DustPuffSmall.png',
}
template = (ROOT/'Assets/PortfolioFX/Textures/NoiseRG.png.meta').read_text(encoding='utf-8-sig')
for name, source in sources.items():
    dest = kit/name
    if dest.exists():
        assert sha256(dest.read_bytes()).digest() == sha256(source.read_bytes()).digest(), str(dest)
    else:
        shutil.copy2(source, dest)
    meta = Path(str(dest)+'.meta')
    if not meta.exists():
        text = re.sub(r'^guid: \w+', 'guid: '+uuid.uuid4().hex, template, flags=re.M)
        text = text.replace('enableMipMap: 1','enableMipMap: 0')
        text = re.sub(r'(wrap[UVW]: )0', r'\g<1>1', text)
        meta.write_text(text, encoding='utf-8')
shutil.copy2(kenney/'LICENSE.txt', kit/'Kenney_LICENSE.txt')
(WORK/'asset-sources.json').write_text(json.dumps({n:str(p) for n,p in sources.items()},indent=2),encoding='utf-8')

doc = Document(DOC)
body = doc._element.body
def txt(el):
    return ''.join(el.itertext()) if False else ''.join(el.xpath('.//w:t/text()'))
original = list(body)
start = next(i for i,e in enumerate(original) if txt(e)=='16 SG_GroundSlash 弧形刀光')
end = next(i for i in range(start+1,len(original)) if txt(original[i]).startswith('17 SG_MeteorSurface'))
before = [etree.tostring(e) for e in original[:start]]
after = [etree.tostring(e) for e in original[end:]]
anchor = original[end]
for e in original[start:end]: body.remove(e)
inserted = []

def move(element):
    anchor.addprevious(element)
    inserted.append(element)

def font(run, size=None, code=False):
    run.font.color.rgb=RGBColor(0,0,0)
    if size: run.font.size=Pt(size)
    run.font.name='Consolas' if code else '微软雅黑'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')

def para(text='', style=None, code=False):
    p=doc.add_paragraph(style=style)
    r=p.add_run(text); font(r,8.5 if code else None,code)
    if code:
        p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.line_spacing=Pt(11)
        p.paragraph_format.keep_with_next=False
        p.paragraph_format.keep_together=False
    else:
        p.paragraph_format.keep_with_next=bool(style and style.startswith('Heading'))
        if not style:
            p.paragraph_format.space_after=Pt(6)
            p.paragraph_format.line_spacing=1.15
    move(p._p)
    return p

def table(rows):
    n=len(rows[0]); t=doc.add_table(rows=0,cols=n)
    t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    usable=6.94
    widths=([1.48,2.65,2.81] if n==3 else [1.35,1.10,1.60,2.89])
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge)
        for k,v in {'val':'single','sz':'4','color':'D9D9D9'}.items(): e.set(qn('w:'+k),v)
        borders.append(e)
    props.append(borders)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        trPr=t.rows[-1]._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        if i==0: trPr.append(OxmlElement('w:tblHeader'))
        for j,(cell,text) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
            p.paragraph_format.line_spacing=1.05
            p.paragraph_format.keep_with_next=(i==0)
            r=p.add_run(text);font(r,9);r.bold=(i==0)
            tcPr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
            tcPr.append(margins)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EDF0');tcPr.append(shade)
    move(t._tbl)
    para('')

lines=(ROOT/'Tools/GroundSlashLesson/chapter16.md').read_text(encoding='utf-8').splitlines()
script_line_counts = {}
i=0
while i<len(lines):
    line=lines[i];i+=1
    if not line.strip(): continue
    if line.startswith('|'):
        batch=[line]
        while i<len(lines) and lines[i].startswith('|'):
            batch.append(lines[i]);i+=1
        rows=[[s.strip() for s in l.strip('|').split('|')] for l in batch if not re.match(r'^\|[\s:|\-]+$',l)]
        table(rows)
    elif line.startswith('#'):
        m=re.match(r'^(#+) (.*)$',line);level=len(m[1]);p=para(m[2],'Heading '+str(level))
        if level==1:p.paragraph_format.page_break_before=True
    elif line.startswith('!['):
        match=re.match(r'!\[(.*?)\]\((.*?)\)',line)
        p=doc.add_paragraph();r=p.add_run();r.add_picture(str(ROOT/match[2]),width=Inches(5.3))
        p.alignment=1
        for node in p._p.xpath('.//wp:docPr'):node.set('descr',match[1])
        move(p._p)
    elif line in ('{{SCRIPT}}', '{{PARTICLES_SCRIPT}}', '{{EDITOR_SCRIPT}}'):
        script_paths = {
            '{{SCRIPT}}': 'Runtime/FXGroundSlashLesson.cs',
            '{{PARTICLES_SCRIPT}}': 'Runtime/FXGroundSlashParticles.cs',
            '{{EDITOR_SCRIPT}}': 'Editor/FXGroundSlashLessonEditor.cs',
        }
        script_path = ROOT/'Assets/PortfolioFX/Scripts'/script_paths[line]
        code=script_path.read_text(encoding='utf-8')
        script_line_counts[script_path.name] = len(code.splitlines())
        for sourceLine in code.splitlines():para(sourceLine,code=True)
        para('代码到此结束。以上全部内容属于 '+script_path.name+' 文件。')
    else:para(line)

# Start the following original chapter on a fresh page, without modifying its XML.
sep=para('')
sep.add_run().add_break(__import__('docx.enum.text',fromlist=['WD_BREAK']).WD_BREAK.PAGE)

body_now=list(body)
s=body_now.index(inserted[0]);e=body_now.index(anchor)
assert [etree.tostring(x) for x in body_now[:s]]==before
assert [etree.tostring(x) for x in body_now[e:]]==after

staging=WORK/'authored.docx'
doc.save(staging)
# Keep untouched package parts byte-identical. Only the body, new picture and its relationship change.
with ZipFile(BACKUP/DOC.name) as old, ZipFile(staging) as new:
    changed={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
    new_names=set(new.namelist())-set(old.namelist())
    with ZipFile(WORK/'updated.docx','w',ZIP_DEFLATED) as out:
        for info in old.infolist():
            out.writestr(info,new.read(info.filename) if info.filename in changed else old.read(info.filename))
        for name in sorted(new_names):out.writestr(name,new.read(name))
shutil.copy2(WORK/'updated.docx',DOC)

# A chapter-only copy keeps visual QA bounded while using the original page setup and styles.
qa=Document(DOC)
qb=qa._element.body
nodes=list(qb)
a=next(i for i,x in enumerate(nodes) if txt(x)=='16 SG_GroundSlash 弧形刀光')
b=next(i for i in range(a+1,len(nodes)) if txt(nodes[i]).startswith('17 SG_MeteorSurface'))
for j,x in enumerate(nodes):
    if x.tag==qn('w:sectPr'):continue
    if j<a or j>=b:qb.remove(x)
last=qb[-2]
if last.xpath('.//w:br'):qb.remove(last)
qa.save(WORK/'chapter16-preview.docx')
report={'document':str(DOC),'backup':str(BACKUP/DOC.name),'inserted_blocks':len(inserted),
        'unchanged_prefix_blocks':len(before),'unchanged_suffix_blocks':len(after),
        'source_code_lines':script_line_counts,'assets':list(sources)}
(WORK/'edit-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
